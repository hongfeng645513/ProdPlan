#!/usr/bin/env python3
"""
Convert Machine/Machines.xlsx into src/data/machines.json for the ProdPlan web app.

Each worksheet in the workbook describes one furnace:

    Name              | Furnace 1
    Model             | Vertical 500
    Function          | Graphitization | Carbonization
    Holders           | 2
    GF size           | 300x300
    GF per holder(g)  | 660
    Yield             | 0.9
    Time(hours)       | 0 1 2 3 ...
    Temperature(C)    | 20 1000 1450 ...            (last cell may be "Open")

A "Rules" sheet holds one plain-English scheduling constraint per row, e.g.

    Furnace 1 and Furnace 2 cannot heat at the same time
    It takes one hour to load a furnace before heating
    It takes one hour to unload a furnace after cooling

Those are parsed into machine-readable form for the planner, and the original
sentences are carried through so the app can show operators the rule it applied.

The script derives, per machine:
  * the measured temperature curve (time in hours -> degrees C)
  * the process phases: heating -> hold (soak) -> natural cooling -> unload
  * a Newton's-law-of-cooling fit of the cooling branch, so the (uncertain and
    machine-dependent) cooling tail can be modelled and re-tuned in the web app.

Usage:  python tools/convert_excel.py [path/to/Machines.xlsx] [-o src/data/machines.json]
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  pip install openpyxl")


# --------------------------------------------------------------------------- #
# workbook parsing
# --------------------------------------------------------------------------- #

LABEL_ALIASES = {
    "name": "name",
    "model": "model",
    "function": "function",
    "holders": "holders",
    "gf size": "gfSize",
    "gf per holder(g)": "gfPerHolder",
    "gf per holder (g)": "gfPerHolder",
    "yield": "yield",
}

TIME_LABELS = ("time(hours)", "time (hours)", "time")
TEMP_HINTS = ("temperature", "carbonization t", "graphitization t", "t(co)", "temp")


def norm(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip().lower() if value is not None else ""


def parse_sheet(ws):
    """Return the raw record for one worksheet, or None if it is not a machine."""
    info, times, temps, open_marker = {}, None, None, False

    rows = list(ws.iter_rows(values_only=True))
    for idx, row in enumerate(rows):
        if not row or all(c is None for c in row):
            continue
        label = norm(row[0])
        rest = list(row[1:])

        key = LABEL_ALIASES.get(label)
        if key:
            info[key] = next((c for c in rest if c is not None), None)
            continue

        numeric = [c for c in rest if isinstance(c, (int, float))]

        # the time axis: a long run of consecutive integers starting at 0
        if (label in TIME_LABELS or label == "") and len(numeric) > 5 and numeric[0] == 0:
            if times is None:
                times = numeric
                continue

        # the temperature row: numbers, possibly closed by the string "Open"
        if any(h in label for h in TEMP_HINTS) or (times is not None and temps is None and numeric):
            series = []
            for cell in rest:
                if isinstance(cell, (int, float)):
                    series.append(float(cell))
                elif isinstance(cell, str) and cell.strip().lower().startswith("open"):
                    open_marker = True
                    break
                else:
                    series.append(None)
            # trim trailing blanks
            while series and series[-1] is None:
                series.pop()
            if len(series) > 5:
                temps = series

    if not info.get("name") or times is None or temps is None:
        return None

    points = []
    for i, temperature in enumerate(temps):
        if temperature is None:
            continue
        hour = float(times[i]) if i < len(times) else float(i)
        points.append({"t": hour, "T": float(temperature)})
    points.sort(key=lambda p: p["t"])

    return {
        "id": re.sub(r"[^a-z0-9]+", "-", norm(info["name"])).strip("-"),
        "name": str(info["name"]).strip(),
        "model": str(info.get("model", "")).strip(),
        "function": str(info.get("function", "")).strip(),
        "holders": info.get("holders"),
        "gfSize": str(info.get("gfSize", "")).strip(),
        "gfPerHolder": info.get("gfPerHolder"),
        "yield": info.get("yield"),
        "points": points,
        "hasOpenMarker": open_marker,
    }


# --------------------------------------------------------------------------- #
# the Rules sheet
# --------------------------------------------------------------------------- #

WORD_NUMBERS = {
    "one": 1, "two": 2, "three": 3, "four": 4, "five": 5,
    "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
    "half": 0.5, "an": 1, "a": 1,
}


def slug(text) -> str:
    return re.sub(r"[^a-z0-9]+", "-", norm(text)).strip("-")


def _duration(text):
    """Pull a number of hours out of '... one hour ...' / '... 1.5 hours ...'."""
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:h\b|hour)", text)
    if m:
        return float(m.group(1))
    m = re.search(r"\b(" + "|".join(WORD_NUMBERS) + r")\s+(?:h\b|hour)", text)
    if m:
        return float(WORD_NUMBERS[m.group(1)])
    return None


def parse_rules(ws):
    """Turn the free-text Rules sheet into constraints the planner can apply.

    Anything that isn't recognised is still returned under `raw` (flagged
    `parsed: False`) so a rule can never be silently dropped: the app lists the
    unparsed ones so an operator knows the planner is not enforcing them.
    """
    sentences = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str) and cell.strip():
                sentences.append(cell.strip())

    exclusive, load_h, unload_h, raw = [], None, None, []

    for line in sentences:
        low = norm(line)
        handled = False

        # "Furnace 1 and Furnace 2 cannot heat at the same time"
        if re.search(r"\b(cannot|can not|can't|must not|never)\b", low) and "same time" in low:
            names = re.findall(r"furnace\s*\d+", low)
            if len(names) >= 2:
                exclusive.append({
                    "machines": [slug(n) for n in names],
                    "scope": "power",       # ramp + soak, i.e. while the element is on
                    "text": line,
                })
                handled = True

        # "It takes one hour to load a furnace before heating"
        elif "load" in low and "unload" not in low:
            d = _duration(low)
            if d is not None:
                load_h, handled = d, True
        elif "unload" in low:
            d = _duration(low)
            if d is not None:
                unload_h, handled = d, True

        raw.append({"text": line, "parsed": handled})

    return {
        "raw": raw,
        "exclusiveHeating": exclusive,
        "loadHours": load_h if load_h is not None else 0.0,
        "unloadHours": unload_h if unload_h is not None else 0.0,
    }


# --------------------------------------------------------------------------- #
# phase detection + Newton cooling fit
# --------------------------------------------------------------------------- #

def detect_phases(points):
    """Split the curve into heating / hold (soak) / cooling.

    Heating is the controlled ramp up to the set point; the hold is the plateau
    at (or within 1% of) the peak; everything after is natural cooling.
    """
    temps = [p["T"] for p in points]
    peak = max(temps)
    plateau = [i for i, T in enumerate(temps) if T >= peak * 0.99]
    heat_end_i, hold_end_i = plateau[0], plateau[-1]
    return {
        "peakTemp": peak,
        "heatStart": points[0]["t"],
        "heatEnd": points[heat_end_i]["t"],      # set point reached
        "holdEnd": points[hold_end_i]["t"],      # heating switched off
        "endOfRecord": points[-1]["t"],
        "heatEndIndex": heat_end_i,
        "holdEndIndex": hold_end_i,
    }


def fit_newton(points, phases, ambient):
    """Fit Newton's law of cooling, anchored at the moment heating is switched off.

        T(t) = Tamb + (Toff - Tamb) * exp(-k * (t - tOff))

    Toff and tOff come straight from the curve, so only the cooling constant k
    is fitted (least squares on the log of the excess temperature). One free
    parameter means the app can expose k as a single "how fast does this furnace
    cool" dial that stays anchored to the measured heat-off point.
    """
    t_off = phases["holdEnd"]
    T_off = points[phases["holdEndIndex"]]["T"]
    excess0 = T_off - ambient
    if excess0 <= 0:
        return None

    usable = [p for p in points if p["t"] > t_off and p["T"] - ambient > 1]
    if len(usable) < 2:
        return None

    num = sum((p["t"] - t_off) * -math.log((p["T"] - ambient) / excess0) for p in usable)
    den = sum((p["t"] - t_off) ** 2 for p in usable)
    if den == 0:
        return None
    k = num / den

    def model(t):
        return ambient + excess0 * math.exp(-k * (t - t_off))

    residuals = [p["T"] - model(p["t"]) for p in usable]
    rmse = math.sqrt(sum(r * r for r in residuals) / len(residuals))

    return {
        "k": round(k, 5),
        "ambient": ambient,
        "tOff": t_off,
        "tempAtOff": T_off,
        "rmse": round(rmse, 1),
        "fittedFrom": usable[0]["t"],
        "fittedTo": usable[-1]["t"],
        "nPoints": len(usable),
        "halfLifeHours": round(math.log(2) / k, 2) if k > 0 else None,
    }


def build(machine):
    points = machine.pop("points")
    phases = detect_phases(points)
    ambient = min(points[0]["T"], min(p["T"] for p in points))
    ambient = min(ambient, 25.0)
    cooling = fit_newton(points, phases, ambient)

    # unload: explicit "Open" marker, else the end of the record
    unload_t = points[-1]["t"]
    unload_T = points[-1]["T"]

    holders = machine["holders"] or 0
    per_holder = machine["gfPerHolder"] or 0
    yld = machine["yield"] or 1
    batch_in = holders * per_holder
    batch_out = batch_in * yld

    machine.update(
        {
            "measured": points,
            "phases": {
                "heatStart": phases["heatStart"],
                "heatEnd": phases["heatEnd"],
                "holdEnd": phases["holdEnd"],
                "unloadAt": unload_t,
                "unloadTemp": unload_T,
                "peakTemp": phases["peakTemp"],
                "heatDuration": round(phases["heatEnd"] - phases["heatStart"], 2),
                "holdDuration": round(phases["holdEnd"] - phases["heatEnd"], 2),
                "coolDuration": round(unload_t - phases["holdEnd"], 2),
                "cycleDuration": round(unload_t - phases["heatStart"], 2),
            },
            "cooling": cooling,
            "capacity": {
                "batchInputG": round(batch_in, 1),
                "batchOutputG": round(batch_out, 1),
                "gramsPerHour": round(batch_out / (unload_t - phases["heatStart"]), 1)
                if unload_t > phases["heatStart"]
                else None,
            },
            # hours where the curve is interpolated rather than measured
            "gaps": find_gaps(points),
        }
    )
    return machine


def find_gaps(points, step=1.0):
    gaps = []
    for a, b in zip(points, points[1:]):
        if b["t"] - a["t"] > step * 1.5:
            gaps.append({"from": a["t"], "to": b["t"]})
    return gaps


# --------------------------------------------------------------------------- #

def main():
    root = Path(__file__).resolve().parent.parent
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook", nargs="?", default=str(root / "Machine" / "Machines.xlsx"))
    ap.add_argument("-o", "--out", default=str(root / "src" / "data" / "machines.json"))
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    machines, rules = [], None
    for ws in wb.worksheets:
        raw = parse_sheet(ws)
        if raw is None:
            if norm(ws.title).startswith("rule"):
                rules = parse_rules(ws)
                print(f"  read rules from sheet '{ws.title}'", file=sys.stderr)
            else:
                print(f"  skipped sheet '{ws.title}' (no machine data)", file=sys.stderr)
            continue
        machines.append(build(raw))

    machines.sort(key=lambda m: m["name"])
    if rules is None:
        rules = {"raw": [], "exclusiveHeating": [], "loadHours": 0.0, "unloadHours": 0.0}

    known = {m["id"] for m in machines}
    for group in rules["exclusiveHeating"]:
        missing = [i for i in group["machines"] if i not in known]
        if missing:
            print(f"  WARNING: rule names unknown furnace(s) {missing}: {group['text']}", file=sys.stderr)

    payload = {
        "source": Path(args.workbook).name,
        "machines": machines,
        "rules": rules,
    }

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    print(f"wrote {out} — {len(machines)} machines")
    print(
        f"  rules: load {rules['loadHours']} h, unload {rules['unloadHours']} h, "
        f"{len(rules['exclusiveHeating'])} exclusive-heating group(s)"
    )
    for r in rules["raw"]:
        if not r["parsed"]:
            print(f"  NOT ENFORCED (unrecognised rule): {r['text']}", file=sys.stderr)
    for m in machines:
        c = m["cooling"]
        print(
            f"  {m['name']:<10} {m['function']:<15} peak {m['phases']['peakTemp']:>5.0f}C  "
            f"heat {m['phases']['heatDuration']:>4}h  hold {m['phases']['holdDuration']:>4}h  "
            f"cool {m['phases']['coolDuration']:>4}h  "
            + (f"k={c['k']:.4f}/h rmse={c['rmse']}C" if c else "no cooling fit")
        )


if __name__ == "__main__":
    main()
